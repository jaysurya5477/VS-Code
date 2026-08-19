# -*- coding: utf-8 -*-
"""
Style engine reproducing the ALIMCO Sales Dashboard documentation set
(D:/New/VS Code/UI5 Practice/Sales Dashboard/Documentation/*.docx).

Every constant below was measured off those three files with python-docx,
so output from this module is visually identical to the reference.
"""
import docx
from docx.shared import Pt, Inches, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ---------------------------------------------------------------- palette ----
INK        = "1B345E"   # headings / table header fill
ACCENT     = "5470C6"   # H2 + H1 underline
MUTED      = "6A6D70"   # subtitle / meta
BODY_DARK  = "22262A"   # H3
CALLOUT    = "2E7D32"   # callout bar + label
BAND       = "F4F6F8"   # table zebra band + code background
LABELFILL  = "EEF2F7"   # label-column fill on spec tables
CODE_INK   = "1A1A1A"

# ------------------------------------------------------------------ sizes ----
SZ_TITLE, SZ_SUBTITLE, SZ_ORG, SZ_META = Pt(30), Pt(14), Pt(11), Pt(9.5)
SZ_H1, SZ_H2, SZ_H3 = Pt(17), Pt(13), Pt(11)
SZ_BODY, SZ_CALLOUT, SZ_CODE = Pt(10.5), Pt(10), Pt(8.5)
SZ_TBL_HEAD, SZ_TBL_BODY = Pt(9.5), Pt(9)


def _shade(el, fill):
    """Applies a solid background fill to a <w:p> or <w:tc>."""
    tag = "w:pPr" if el.tag == qn("w:p") else "w:tcPr"
    pr = el.find(qn(tag))
    if pr is None:
        pr = OxmlElement(tag)
        el.insert(0, pr)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pr.append(shd)


def _border(par, side, sz, color, space=0):
    """Adds one paragraph border edge - used for the H1 rule and callout bar."""
    pPr = par._p.get_or_add_pPr()
    bdr = pPr.find(qn("w:pBdr"))
    if bdr is None:
        bdr = OxmlElement("w:pBdr")
        pPr.append(bdr)
    e = OxmlElement("w:" + side)
    e.set(qn("w:val"), "single")
    e.set(qn("w:sz"), str(sz))
    e.set(qn("w:space"), str(space))
    e.set(qn("w:color"), color)
    bdr.append(e)


def _run(par, text, size=None, bold=None, color=None, font=None):
    r = par.add_run(text)
    if size:
        r.font.size = size
    if bold is not None:
        r.bold = bold
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    if font:
        r.font.name = font
        r._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    return r


class Doc:
    """Thin builder over python-docx carrying the reference look."""

    def __init__(self, title, subtitle, meta):
        self.d = docx.Document()

        s = self.d.sections[0]
        s.page_width, s.page_height = Inches(8.5), Inches(11)
        s.left_margin = s.right_margin = Emu(822960)   # 0.9"
        s.top_margin = s.bottom_margin = Emu(731520)   # 0.8"

        n = self.d.styles["Normal"]
        n.font.name = "Calibri"
        n.font.size = SZ_BODY
        n.paragraph_format.space_after = Pt(6)
        n.paragraph_format.line_spacing = 1.15

        for st in ("List Bullet", "List Number"):
            pf = self.d.styles[st].paragraph_format
            pf.space_after = Pt(3)
            pf.left_indent = Inches(0.3)

        self._cover(title, subtitle, meta)

    # ------------------------------------------------------------- cover ----
    def _cover(self, title, subtitle, meta):
        for _ in range(6):
            self.d.add_paragraph()

        def centred(text, size, color, bold=None):
            p = self.d.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _run(p, text, size, bold, color)

        centred(title, SZ_TITLE, INK, True)
        centred(subtitle, SZ_SUBTITLE, MUTED)
        self.d.add_paragraph()
        centred("ALIMCO  \u00b7  SAP S/4HANA  \u00b7  Sales & Distribution", SZ_ORG, ACCENT, True)
        centred(meta, SZ_META, MUTED)
        self.d.add_paragraph()

    # ---------------------------------------------------------- headings ----
    def h1(self, text):
        p = self.d.add_paragraph()
        pf = p.paragraph_format
        pf.space_before, pf.space_after = Pt(16), Pt(8)
        _run(p, text, SZ_H1, True, INK)
        _border(p, "bottom", 8, ACCENT)
        return p

    def h2(self, text):
        p = self.d.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        _run(p, text, SZ_H2, True, ACCENT)
        return p

    def h3(self, text):
        p = self.d.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        _run(p, text, SZ_H3, True, BODY_DARK)
        return p

    # -------------------------------------------------------------- body ----
    def p(self, text):
        par = self.d.add_paragraph()
        _run(par, text)
        return par

    def bullets(self, items):
        for it in items:
            par = self.d.add_paragraph(style="List Bullet")
            _run(par, it)

    def numbers(self, items):
        for it in items:
            par = self.d.add_paragraph(style="List Number")
            _run(par, it)

    def callout(self, label, text):
        """Green left-barred note - the reference's recurring emphasis block."""
        par = self.d.add_paragraph()
        pf = par.paragraph_format
        pf.space_before, pf.space_after = Pt(6), Pt(8)
        pf.left_indent = Inches(0.15)
        _run(par, label + ":  ", SZ_CALLOUT, True, CALLOUT)
        _run(par, text, SZ_CALLOUT)
        _border(par, "left", 18, CALLOUT, space=6)
        return par

    def code(self, text):
        par = self.d.add_paragraph()
        pf = par.paragraph_format
        pf.space_before, pf.space_after = Pt(4), Pt(8)
        pf.left_indent = Inches(0.25)
        pf.line_spacing = 1.0
        _run(par, text, SZ_CODE, None, CODE_INK, "Consolas")
        _shade(par._p, BAND)
        return par

    def spacer(self):
        self.d.add_paragraph()

    # ------------------------------------------------------------ tables ----
    def _grid(self, rows, cols, widths):
        t = self.d.add_table(rows=rows, cols=cols)
        t.style = "Table Grid"
        if widths:
            for r in t.rows:
                for c, w in zip(r.cells, widths):
                    c.width = Inches(w)
        return t

    def table(self, header, rows, widths=None):
        """Header-row table: dark banner row + zebra-banded body."""
        t = self._grid(len(rows) + 1, len(header), widths)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER

        for i, txt in enumerate(header):
            cell = t.rows[0].cells[i]
            cell.text = ""
            _run(cell.paragraphs[0], txt, SZ_TBL_HEAD, True, "FFFFFF")
            _shade(cell._tc, INK)

        for ri, row in enumerate(rows):
            band = (ri % 2 == 1)
            for ci, txt in enumerate(row):
                cell = t.rows[ri + 1].cells[ci]
                cell.text = ""
                _run(cell.paragraphs[0], str(txt), SZ_TBL_BODY)
                if band:
                    _shade(cell._tc, BAND)
        self.spacer()
        return t

    def spec(self, pairs, widths=(2.0, 4.6)):
        """Two-column label/value table: tinted bold label column, no banding."""
        t = self._grid(len(pairs), 2, widths)
        for ri, (k, v) in enumerate(pairs):
            lc = t.rows[ri].cells[0]
            lc.text = ""
            _run(lc.paragraphs[0], k, SZ_TBL_BODY, True)
            _shade(lc._tc, LABELFILL)

            vc = t.rows[ri].cells[1]
            vc.text = ""
            _run(vc.paragraphs[0], str(v), SZ_TBL_BODY)
        self.spacer()
        return t

    # -------------------------------------------------------------- save ----
    def save(self, path):
        self.d.save(path)
        return path
